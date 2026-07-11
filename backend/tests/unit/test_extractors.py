"""Unit tests for per-format text extractors (real bytes built in-test)."""

from __future__ import annotations

from io import BytesIO

import pytest
from app.domain.interfaces.extraction import TextExtractionError
from app.infrastructure.extraction import extractor_for

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docx_bytes(*paragraphs: str) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(text: str) -> bytes:
    """Assemble a minimal valid single-page PDF (correct xref offsets)."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_position = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_position}\n%%EOF\n"
    ).encode()
    return bytes(out)


async def test_plain_text_extraction() -> None:
    text = await extractor_for("text/plain").extract(b"hello\nworld")
    assert text == "hello\nworld"


async def test_csv_uses_plain_text_extractor() -> None:
    text = await extractor_for("text/csv").extract(b"a,b\n1,2\n")
    assert "a,b" in text


async def test_docx_extraction() -> None:
    content = _docx_bytes("First paragraph.", "Second paragraph.")
    text = await extractor_for(_DOCX_MIME).extract(content)
    assert "First paragraph." in text
    assert "Second paragraph." in text


async def test_xlsx_extraction() -> None:
    content = _xlsx_bytes([["name", "amount"], ["widget", 42]])
    text = await extractor_for(_XLSX_MIME).extract(content)
    assert "name\tamount" in text
    assert "widget\t42" in text


async def test_pdf_extraction() -> None:
    text = await extractor_for("application/pdf").extract(_pdf_bytes("Hello PDF world"))
    assert "Hello PDF world" in text


async def test_pdf_without_text_raises() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(TextExtractionError, match="no extractable text"):
        await extractor_for("application/pdf").extract(buffer.getvalue())


async def test_malformed_docx_raises() -> None:
    with pytest.raises(TextExtractionError):
        await extractor_for(_DOCX_MIME).extract(b"PK\x03\x04 not really a docx")


def test_unknown_mime_type_raises() -> None:
    with pytest.raises(TextExtractionError):
        extractor_for("application/x-unknown")
